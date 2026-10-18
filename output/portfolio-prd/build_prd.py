from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path(__file__).resolve().parent / "Miebaka Portfolio Product Requirements Document.docx"

NAVY = "101828"
BLUE = "2643E9"
PALE = "EEF2FF"
GREY = "667085"
LIGHT = "EAECF0"
WHITE = "FFFFFF"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margin(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def borderless(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "nil")


def keep_with_next(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext")
    pPr.append(keep)


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.62)
sec.bottom_margin = Inches(0.62)
sec.left_margin = Inches(0.72)
sec.right_margin = Inches(0.72)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(9.2)
styles["Normal"].font.color.rgb = RGBColor.from_string(NAVY)
styles["Normal"].paragraph_format.space_after = Pt(4.5)
styles["Normal"].paragraph_format.line_spacing = 1.05

for s, size, color, before, after in [
    ("Title", 35, NAVY, 0, 12),
    ("Subtitle", 14, GREY, 0, 10),
    ("Heading 1", 20, NAVY, 15, 6),
    ("Heading 2", 13, BLUE, 10, 4),
    ("Heading 3", 10.5, NAVY, 7, 2),
]:
    st = styles[s]
    st.font.name = "Aptos Display" if s != "Normal" else "Aptos"
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string(color)
    st.font.bold = s in ("Title", "Heading 1", "Heading 2", "Heading 3")
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

if "PRD Label" not in styles:
    label = styles.add_style("PRD Label", WD_STYLE_TYPE.PARAGRAPH)
    label.font.name = "Aptos"
    label.font.size = Pt(7.5)
    label.font.bold = True
    label.font.color.rgb = RGBColor.from_string(BLUE)
    label.paragraph_format.space_after = Pt(3)

for section in doc.sections:
    header = section.header
    p = header.paragraphs[0]
    p.text = "MIEBAKA JOSHUA  /  PORTFOLIO PRD"
    p.style = styles["PRD Label"]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("13 OCT 2026   •   ")
    r.font.name = "Aptos"
    r.font.size = Pt(7.5)
    r.font.color.rgb = RGBColor.from_string(GREY)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)


def label(text):
    p = doc.add_paragraph(style="PRD Label")
    p.add_run(text.upper())
    return p


def para(text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def bullets(items, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(2.2)
        p.add_run(item)


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.style = "Table Grid"
    set_repeat_table_header(t.rows[0])
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = h.upper()
        shade(c, NAVY)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margin(c)
        for r in c.paragraphs[0].runs:
            r.font.bold = True
            r.font.size = Pt(7.4)
            r.font.color.rgb = RGBColor.from_string(WHITE)
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margin(cells[i])
            if ridx % 2 == 1:
                shade(cells[i], "F8FAFC")
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.size = Pt(8)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


# Cover
doc.add_paragraph().paragraph_format.space_after = Pt(78)
label("Product requirements document")
p = doc.add_paragraph(style="Title")
p.add_run("Miebaka Portfolio")
p = doc.add_paragraph(style="Subtitle")
p.add_run("A proof-led structural upgrade that preserves the restored visual character")
doc.add_paragraph().paragraph_format.space_after = Pt(22)
table(
    ["Document", "Value"],
    [
        ("Owner", "Miebaka Joshua"),
        ("Status", "Implementation-ready PRD"),
        ("Product", "miebaka.design portfolio"),
        ("Visual authority", "Restored production baseline — commit 54dcc41"),
        ("Source", "Portfolio Design Engineering Review and implementation plan"),
        ("Date", "13 October 2026"),
    ],
    [1.55, 5.15],
)
doc.add_paragraph().paragraph_format.space_after = Pt(30)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = p.add_run("PRODUCT DECISION")
r.bold = True; r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string(BLUE)
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(0.26)
p.paragraph_format.right_indent = Inches(0.6)
r = p.add_run("Improve structure, proof, accessibility and conversion without redesigning the site’s visual identity. The restored production website is the visual source of truth; the proof-sheet contributes selective editorial mechanisms, not a replacement skin.")
r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = RGBColor.from_string(NAVY)
doc.add_page_break()

# 1
doc.add_heading("1. Executive summary", level=1)
para("The portfolio already has a distinctive visual voice: editorial restraint, an Ultramarine–paper–ink palette, sharp geometry, authored imagery, quiet motion, a strong portrait treatment and warm, direct writing. Its primary weakness is not visual taste; it is insufficient proof. Hiring managers and potential clients can see an authored surface, but cannot consistently reach complete case studies, evaluate the work, submit a reliable enquiry or understand the full range of practice.")
para("This release turns the current site into a credible senior-design portfolio by making the proof layer complete and trustworthy. It opens one flagship case study first, restructures the homepage around evidence, adds truthful Explorations and Community destinations, strengthens the Index, ships a real contact service, makes the CV accessible, and introduces robust accessibility, performance and content-governance requirements.")
para("The release must feel like the same website, only more resolved. No requirement authorizes a stylistic reset. Proof-sheet ideas may be used only where they clarify production, sequencing or authorship and remain subordinate to the restored visual system.")

doc.add_heading("2. Product vision and positioning", level=1)
doc.add_heading("2.1 Vision", level=2)
para("Create a precise, memorable portfolio that lets design leaders quickly understand Miebaka’s point of view, inspect the quality of his decisions and production thinking, and confidently begin a conversation.")
doc.add_heading("2.2 Positioning", level=2)
para("Miebaka is presented as a senior multidisciplinary designer whose value lies in the connection between strategic thinking, visual systems, editorial judgment and production discipline—not as a generalist offering an undifferentiated list of services.")
doc.add_heading("2.3 Primary product job", level=2)
para("Help hiring managers and design leaders evaluate Miebaka for senior design employment. Client enquiries and professional reputation are important secondary jobs.")
doc.add_heading("2.4 Target roles and markets", level=2)
bullets([
    "Primary roles: Art Director, Senior Brand Designer and Senior Visual Designer.",
    "Markets: remote international roles, Nigerian roles and relocation opportunities.",
    "Secondary opportunities: selected identity, campaign, editorial and creative-direction engagements.",
])

doc.add_heading("3. Objectives and success definition", level=1)
table(["Priority", "Objective", "Success signal"], [
    ("P0", "Make professional work evaluable", "At least the TEDx RSU flagship case study is public, complete and navigable."),
    ("P0", "Make contact reliable", "The form sends through a server endpoint, reports clear states and preserves input after recoverable failure."),
    ("P0", "Preserve exact visual character", "Owner review confirms continuity with the restored baseline; no proof-sheet or generic UI drift."),
    ("P1", "Clarify the breadth of practice", "Work, Explorations, Community, About, Services and Writing have truthful destinations and states."),
    ("P1", "Strengthen professional trust", "CV, verified metadata, clear roles and factual case-study captions are available."),
    ("P1", "Improve inclusive quality", "Keyboard, reduced-motion, no-JS, 320px, zoom and assistive-technology release checks pass."),
    ("P1", "Protect speed and stability", "Core Web Vitals and bundle budgets meet the release gates in this PRD."),
], [0.55, 2.35, 3.85])
doc.add_heading("3.1 Conversion hierarchy", level=2)
bullets([
    "Primary conversion: send a credible contact enquiry.",
    "Secondary conversions: complete a case study, download the CV, visit LinkedIn, remember the work and return.",
    "Micro-conversions: open the Index, inspect services, view project metadata and continue to the next project.",
])

doc.add_heading("4. Users and priority journeys", level=1)
table(["Audience", "Need", "Required journey"], [
    ("Design leader / hiring manager", "Fast evidence of senior judgment, role clarity and craft", "Home → selected work → TEDx RSU → About/CV → Contact"),
    ("Founder / client", "Confidence in fit, process and delivery", "Home → Services → relevant project → Contact"),
    ("Creative director / collaborator", "A clear view of authorship and complementary strengths", "Index → Work/Explorations → project detail → Contact"),
    ("General visitor", "A legible introduction without specialist knowledge", "Home → selected work/profile → Index"),
    ("Designer / peer", "Process, production and editorial detail", "Work → project plates/captions → Writing"),
], [1.35, 2.2, 3.2])
doc.add_heading("4.1 Journey requirements", level=2)
bullets([
    "Every route must make the next meaningful action clear without adding loud, conversion-template styling.",
    "The visitor must never encounter a locked professional project without a truthful explanation and an available published alternative.",
    "The full site must remain understandable with JavaScript disabled; enhancements may enrich, not reveal, core content.",
])

doc.add_heading("5. Current-state problem statement", level=1)
table(["Problem", "Impact", "Priority"], [
    ("All professional case studies are locked", "The site asks for trust without supplying enough evidence.", "P0"),
    ("Homepage biography precedes proof", "The strongest hiring-manager journey is delayed.", "P0"),
    ("Explorations and Community are missing", "The navigation cannot truthfully represent the full practice.", "P1"),
    ("Contact is mailto-only and the API route fails", "Primary conversion is unreliable and hard to measure.", "P0"),
    ("Tests protect an obsolete contract", "Green checks can validate the wrong product behavior.", "P0"),
    ("Navigation and CV access are incomplete", "Users must infer routes and professional history.", "P1"),
    ("Client loop lacks a visible pause control", "Motion is not fully operable or inclusive.", "P1"),
    ("Content, CSS and JavaScript are overly consolidated", "Changes carry avoidable regression and maintenance risk.", "P1"),
    ("Some claims lack evidence", "Professional trust can be weakened by unsupported outcomes.", "P0"),
], [2.25, 3.75, 0.7])

doc.add_heading("6. Product and design principles", level=1)
bullets([
    "Preserve before extending. The restored production site at commit 54dcc41 is the visual authority.",
    "Proof before biography. Show the work early, then deepen the human story.",
    "Truth over completeness. A candid coming-soon or empty state is better than invented work, metrics or testimony.",
    "One visual argument. New UI must belong to the existing editorial system, not introduce cards, pills or dashboard conventions.",
    "Mechanism must reveal meaning. Expressive interactions are allowed only when they reveal process, material, sequencing or authorship.",
    "Content before animation. Essential content, actions and navigation must exist in the document before enhancement runs.",
    "Quiet confidence. Motion, labels and controls should feel precise and deliberate rather than theatrical.",
])
doc.add_heading("6.1 Immutable visual characteristics", level=2)
bullets([
    "Ultramarine, paper and ink remain the governing neutral/accent relationship; authored project colours remain project-specific.",
    "Retain the current typography hierarchy and three-typeface ceiling.",
    "Retain low density, sharp geometry, asymmetric air, quiet motion, portrait treatment and signature/footer character.",
    "Do not replace the site with the proof-sheet interface, Material 3, a generic portfolio template or a component-library aesthetic.",
])

doc.add_heading("7. Scope", level=1)
doc.add_heading("7.1 In scope", level=2)
bullets([
    "Reordered homepage with selected work before profile, followed by services, client proof and contact.",
    "Mixed publication states for projects, with one complete flagship case study at launch.",
    "New Explorations and Community destinations with supplied entries or truthful empty states.",
    "Expanded full-screen editorial Index containing all routes, contact, CV and preference controls.",
    "Downloadable CV linked from About and Index.",
    "Five-service swatch interaction, experience ledger and controlled client-logo loop.",
    "Server-backed contact form, updated privacy content, structured error handling and abuse protection.",
    "Content-model split, build-time validation, CSS/JS consolidation and production-output cleanup.",
    "Accessibility, performance, browser, security and visual-regression release gates.",
])
doc.add_heading("7.2 Explicit non-goals", level=2)
bullets([
    "Recreating the rejected thesis-aligned redesign or replacing the current visual language.",
    "Publishing incomplete projects as finished work or inventing outcomes, testimonials, communities or principles.",
    "Adding authentication, a database, a CMS, advertising trackers or a large analytics stack.",
    "Making every proof-sheet motif global, adopting a fourth typeface, a custom cursor or mandatory smooth-scroll libraries.",
    "Deploying material visual changes without preview review and explicit owner approval.",
])

doc.add_heading("8. Information architecture", level=1)
table(["Route", "Purpose", "Launch state"], [
    ("Home", "Positioning, selected proof, profile, services, trust and contact", "Required"),
    ("Work", "Published and forthcoming professional projects", "Required"),
    ("Work / TEDx RSU", "Flagship end-to-end case study", "Required"),
    ("Explorations", "Self-initiated, experimental or process-led work", "Required; truthful empty state allowed"),
    ("Community", "Community activity, contribution and convening", "Required; truthful empty state allowed"),
    ("About", "Biography, experience, capabilities and CV access", "Required"),
    ("Services", "Five clear service offers and engagement fit", "Required"),
    ("Writing", "Existing essays and future authored thinking", "Retain"),
    ("Contact", "Direct email and reliable form submission", "Required"),
    ("Privacy", "Contact-data and site-privacy disclosure", "Footer only"),
], [1.3, 3.5, 1.85])
doc.add_heading("8.1 Navigation model", level=2)
bullets([
    "Keep the desktop navigation persistent and visually quiet.",
    "Use the Index as the comprehensive navigation layer; it must include every route, contact details, CV and theme/motion preferences.",
    "Do not introduce a bottom navigation bar unless testing demonstrates a clear benefit without compromising the visual character.",
])

doc.add_heading("9. Functional requirements by surface", level=1)
doc.add_heading("9.1 Homepage", level=2)
table(["Order", "Module", "Requirement"], [
    ("1", "Hero", "Retain the restored composition and positioning. Add a support line only if it does not weaken the hierarchy."),
    ("2", "Selected work", "Show published work first. Cards/links must disclose publication state and be keyboard-operable."),
    ("3", "Profile / portrait", "Retain the quiet photographic treatment; do not restore the retired halftone portrait."),
    ("4", "Services", "Show five offers through the swatch mechanism, with all content available without motion."),
    ("5", "Client proof", "Use factual client marks/names; the loop must have visible pause/play and become static under reduced motion."),
    ("6", "Working principles", "Include only after owner-supplied copy is approved; omit rather than invent."),
    ("7", "Contact", "Provide direct email plus the server-backed form and clear response states."),
    ("8", "Signature/footer", "Preserve the restored signature and footer character."),
], [0.45, 1.45, 4.75])
doc.add_heading("9.2 Work index", level=2)
bullets([
    "Support published, coming-soon and private/archive states through explicit metadata.",
    "Published projects open directly. Coming-soon projects remain visible only when their presence adds useful context and their state is unambiguous.",
    "Do not lock every project. A visitor must always have a complete published case study to evaluate.",
])
doc.add_heading("9.3 Flagship case study: TEDx RSU", level=2)
bullets([
    "Open with a concise proposition and a strong hero asset.",
    "State role, collaborators, year and deliverables close to the opening.",
    "Tell the sequence: context → problem → relevant process → strategic decision → system → applications → production/handoff → verified outcomes → reflection → next project.",
    "Show only process that survived into or materially informed the final direction; do not pad the page with rejected routes.",
    "Use proof-sheet mechanisms selectively: PLATE numbering, factual captions, a restrained chapter/folio system and production details.",
    "Omit metrics or outcomes that cannot be evidenced. Clearly distinguish individual contribution from team contribution.",
])
doc.add_heading("9.4 Explorations and Community", level=2)
bullets([
    "Each route must exist at launch and explain its purpose in plain language.",
    "If approved entries are unavailable, show an intentional empty state with no fabricated dates, scale, partners or outcomes.",
    "Empty-state pages must still include global navigation, Index access and a relevant next action.",
])
doc.add_heading("9.5 About, CV and experience", level=2)
bullets([
    "Present biography, capabilities and experience without duplicating the homepage word-for-word.",
    "Use an experience ledger for role, organisation, period and concise contribution; treat it as evidence, not decoration.",
    "Offer an approved, accessible PDF CV from About and Index. The CV is not a primary top-level route.",
])
doc.add_heading("9.6 Services", level=2)
bullets([
    "Present five discrete offers with clear problems solved, typical outputs and engagement fit.",
    "Desktop may use a layered swatch fan; mobile must use an accessible horizontal or stacked presentation with no hidden content.",
    "Selection, focus and expanded state must be perceivable without relying on colour or hover alone.",
])
doc.add_heading("9.7 Contact", level=2)
bullets([
    "Require name, valid email and message; apply documented field and request limits.",
    "Show idle, focused, validating, submitting, success, rate-limited, unavailable and recoverable-error states.",
    "Preserve entered values when submission fails. Never clear the form before confirmed success.",
    "Always show a copyable email address and a truthful no-JavaScript fallback.",
])

doc.add_heading("10. Content and evidence requirements", level=1)
table(["Content type", "Rule", "Owner dependency"], [
    ("Project claims", "Publish only verified outcomes, dates, roles and collaborators.", "Approved facts and sources"),
    ("Project media", "Provide final, rights-cleared assets with dimensions, alt intent and captions.", "TEDx RSU media package"),
    ("Explorations", "Use real self-initiated work or a truthful empty state.", "Approved entries"),
    ("Community", "Use documented participation/contribution or a truthful empty state.", "Approved entries"),
    ("Working principles", "Do not infer principles from tone or prior copy.", "Owner-written/approved copy"),
    ("CV", "Use a current, approved and accessible PDF.", "Final resume file"),
    ("Client proof", "Use authorised names/marks; do not imply unverified endorsement.", "Approved client list/assets"),
], [1.35, 3.3, 2.05])
doc.add_heading("10.1 Content governance", level=2)
bullets([
    "Every project record must have a unique slug, publication state, route owner, required metadata and valid next-project reference.",
    "Every media item must declare dimensions, purpose, accessible alternative and loading priority.",
    "A build-time validator must fail on duplicate slugs, invalid routes, contradictory publication states, missing required metadata and broken project relationships.",
    "Editorial language must remain warm, precise and direct. Avoid inflated strategic language, generic design clichés and unsupported numerical claims.",
])

doc.add_heading("11. Interaction, motion and preferences", level=1)
bullets([
    "Motion must explain hierarchy, continuity, state or production logic; decorative motion without a communicative job is out of scope.",
    "The Index, service swatches and project transitions must remain fully usable by keyboard and touch.",
    "Reduced-motion behavior must be designed per component: replace loops and large transforms with static or low-motion alternatives rather than applying a blanket 1ms override.",
    "Theme and motion preferences must persist without blocking first render or exposing a flash that materially changes the intended composition.",
    "Avoid mandatory GSAP, Lenis, ScrollTrigger or SplitText dependencies. Add a library only after the native implementation proves insufficient and the budget still passes.",
    "Do not use a custom cursor, hover-only instructions, cloned semantic text or live decorative feeds that fail silently.",
])

doc.add_heading("12. Accessibility requirements", level=1)
para("WCAG 2.2 AA is the release floor. The product should retain stronger contrast where it already exists, but must not impose a universal 7:1 target when that would damage the approved visual character without meaningful benefit.")
table(["Area", "Requirement"], [
    ("Keyboard", "All routes, controls, overlays, swatches, project links and form states are reachable, visible and operable."),
    ("Focus", "Focus treatment is obvious against every theme and project colour; focus is managed when Index/overlays open and close."),
    ("Touch", "Interactive hit regions target approximately 44×44 CSS pixels where practical."),
    ("Motion", "Pause is available for moving content; reduced-motion alternatives are component-specific and meaningful."),
    ("Forms", "Labels, descriptions, required state and errors are programmatically associated and announced."),
    ("Structure", "Landmarks, heading order, link purpose and current-page state remain coherent across routes."),
    ("Resilience", "Core content and navigation work without JavaScript; verify forced colours, 200% zoom, increased text spacing and 320px width."),
    ("Media", "Alt text communicates purpose; decorative assets are ignored; captions carry factual context rather than repeating the image."),
], [1.25, 5.45])

doc.add_heading("13. Performance and compatibility", level=1)
table(["Metric", "Target", "Hard gate"], [
    ("Largest Contentful Paint", "≤ 2.0 s", "≤ 2.5 s"),
    ("Interaction to Next Paint", "≤ 150 ms", "≤ 200 ms"),
    ("Cumulative Layout Shift", "≤ 0.05", "≤ 0.05"),
    ("Initial compressed JavaScript", "≤ 60 KB", "≤ 75 KB"),
    ("Initial compressed CSS", "≤ 35 KB", "≤ 45 KB"),
    ("Homepage media before interaction", "≤ 1.2 MB", "≤ 1.5 MB"),
    ("Navigation long task", "None over 100 ms", "None over 200 ms"),
], [2.4, 1.85, 2.45])
bullets([
    "Prerender essential copy, navigation and project metadata. Defer non-essential media and enhancements.",
    "Use responsive images, explicit dimensions and appropriate modern formats while preserving authored crops.",
    "Apply immutable caching only to content-hashed assets. Do not cache mutable HTML as immutable.",
    "Respect Save-Data and reduced-motion preferences with quieter media and fewer background requests.",
    "Support the latest two stable versions of Chrome, Safari, Firefox and Edge, plus current iOS Safari and Android Chrome.",
])

doc.add_heading("14. Technical architecture", level=1)
doc.add_heading("14.1 Content model", level=2)
para("Split the current monolithic site data into explicit domains: site configuration, navigation, projects, case studies, services, writing and clients. Keep rendered content as the canonical accessible layer; JavaScript manages enhancement and state rather than becoming the only source of content.")
doc.add_heading("14.2 Frontend boundaries", level=2)
bullets([
    "Create stable modules for navigation/Index, theme, motion preferences, project states, service interaction, marquee control and form state.",
    "Consolidate duplicated CSS into tokens and component-level rules while preserving measured geometry and visual output.",
    "Keep route metadata, canonical URLs, social cards and structured data aligned with publication state.",
    "Treat existing tests as evidence, not authority. Rewrite assertions that encode obsolete locked-project or contact behavior.",
])
doc.add_heading("14.3 Contact service", level=2)
para("Required flow: HTTP handler → request validator → abuse guard → mail-service interface → Resend adapter → structured response and privacy-safe logger.")
table(["Layer", "Requirement"], [
    ("Request", "Same-origin JSON POST; require name/email/message; enforce field/body limits; reject unknown/malformed input."),
    ("Abuse controls", "Honeypot, minimum-fill-time check and shared infrastructure-backed rate limiting—not process memory."),
    ("Provider", "Resend behind a replaceable interface with an 8-second timeout and mapped provider errors."),
    ("Responses", "200/202 success; 400/422 validation; 429 with Retry-After; 503 provider unavailable; 500 unexpected."),
    ("Logging", "Record request ID, status, latency and provider error class; never log message bodies or full email addresses."),
    ("Client", "Maintain values on failure, prevent duplicate submission and provide a direct email fallback."),
], [1.35, 5.35])

doc.add_heading("15. Security and privacy", level=1)
bullets([
    "Preserve and verify existing security headers. Introduce CSP changes in report-only mode before enforcement.",
    "Use Vercel as the canonical runtime and document required environment variables without committing secrets.",
    "Restrict and document any public third-party key. If Last.fm remains, consider a proxy/cache and ensure the portfolio still works when the feed fails.",
    "Ship production-only output: no source maps, internal drafts, private notes, test fixtures or unused raw assets.",
    "Do not add advertising trackers. If measurement is introduced, minimise collection and update the privacy page before release.",
    "Update privacy content to explain contact submissions, retention, processors and a clear deletion/contact route.",
])

doc.add_heading("16. Measurement plan", level=1)
para("V1 success is judged primarily through product completeness, quality gates and direct conversion reliability. A new analytics platform is not required for launch. If privacy-conscious analytics already exists or is later approved, measure only events needed to improve the priority journeys.")
table(["Signal", "Definition", "Use"], [
    ("Case-study completion", "Visitor reaches the next-project/reflection region of TEDx RSU", "Assess whether proof is being consumed"),
    ("Contact success rate", "Accepted submissions divided by valid attempts", "Monitor conversion reliability"),
    ("Contact error classes", "Validation, rate-limit, provider and unexpected failures", "Identify operational defects without storing message content"),
    ("CV downloads", "Approved CV download action", "Assess professional-intent pathway"),
    ("Index usage", "Open and route selection", "Assess wayfinding"),
    ("Core Web Vitals", "Field and controlled-test LCP/INP/CLS", "Protect perceived quality"),
], [1.55, 3.15, 2.0])

doc.add_heading("17. Dependencies and owner-supplied inputs", level=1)
table(["Dependency", "Needed for", "Fallback"], [
    ("TEDx RSU assets and verified facts", "Flagship case study", "Do not publish incomplete or speculative content"),
    ("Explorations entries", "Explorations route", "Approved truthful empty state"),
    ("Community entries", "Community route", "Approved truthful empty state"),
    ("Current accessible CV PDF", "About and Index", "Hide download action until supplied"),
    ("Verified metrics/collaborators", "Project evidence", "Omit the claim or collaborator"),
    ("Working-principles copy", "Optional homepage module", "Omit the module"),
    ("Vercel and Resend configuration", "Contact delivery", "Direct email remains visible; form cannot ship as complete"),
    ("Shared rate-limit provider", "Abuse prevention", "Do not substitute in-memory production limits"),
    ("Owner preview approval", "Production release", "Hold deployment"),
], [2.0, 2.75, 1.95])

doc.add_heading("18. Risks and mitigations", level=1)
table(["Risk", "Mitigation", "Release consequence"], [
    ("Visual drift toward the proof sheet or rejected build", "Compare every changed route with the restored baseline; require owner visual approval.", "Block release"),
    ("Scope expands into a redesign", "Use the in-scope/non-goal boundary and prioritise proof, routes and reliability.", "Defer nonessential concepts"),
    ("Missing content encourages invention", "Use explicit content dependencies and truthful empty/omitted states.", "Ship less, never fabricate"),
    ("Animation causes accessibility or performance regressions", "Progressive enhancement, per-component reduced motion and budget tests.", "Remove or simplify motion"),
    ("Contact provider or abuse controls fail", "Adapter boundary, timeout, structured errors, shared limit and fallback email.", "Block form-complete claim"),
    ("Tests validate old behavior", "Rewrite contract tests before feature implementation; add route and state coverage.", "Block release"),
    ("Heavy media undermines perceived quality", "Responsive media pipeline, explicit dimensions, preload only the real LCP asset.", "Optimise or defer assets"),
], [2.15, 3.35, 1.2])

doc.add_heading("19. Delivery plan", level=1)
table(["Phase", "Outcome", "Exit criterion"], [
    ("0 — Protect baseline", "Snapshot visual/technical state; rewrite obsolete expectations", "Baseline review set and current contract documented"),
    ("1 — Model and IA", "Split content domains; add route/publication schema and validator", "Build fails correctly on invalid content"),
    ("2 — Home and navigation", "Proof-first homepage, strengthened Index and truthful project states", "All primary routes reachable and keyboard-operable"),
    ("3 — Flagship proof", "Complete TEDx RSU case study", "Content, media, captions and claims approved"),
    ("4 — Professional evidence", "About, experience ledger, services and CV", "No unsupported claims; CV works"),
    ("5 — Contact and privacy", "Server contact, form states, abuse controls and privacy update", "End-to-end delivery and failure tests pass"),
    ("6 — System consolidation", "CSS/JS boundaries, motion and responsive behavior", "No visual regression against baseline"),
    ("7 — Release QA", "Accessibility, performance, security, browser and no-JS checks", "All hard gates pass"),
    ("8 — Preview and release", "Owner-reviewed preview and production deployment", "Explicit owner approval and post-deploy smoke test"),
], [1.2, 3.15, 2.35])

doc.add_heading("20. Acceptance criteria and release gates", level=1)
doc.add_heading("20.1 Product completeness", level=2)
bullets([
    "TEDx RSU is public, complete and reachable from Home and Work.",
    "Every project link has a truthful destination or explicit state; no false locks or dead-end cards.",
    "Explorations, Community, About, Services, Writing, Contact and Privacy resolve without broken assets.",
    "CV download works from About and Index when the approved file is supplied.",
    "The contact form delivers through the endpoint and correctly handles validation, duplicate, rate-limit, provider and unexpected failures.",
])
doc.add_heading("20.2 Quality gates", level=2)
bullets([
    "Owner approves the preview as visually continuous with the restored production baseline.",
    "No unsupported metrics, project outcomes, client relationships or collaborator credits remain.",
    "Keyboard, focus, screen-reader basics, reduced motion, forced colours, 200% zoom, increased spacing and 320px tests pass.",
    "Core navigation, project content and contact fallback work without JavaScript.",
    "Performance budgets and hard Core Web Vitals gates pass on representative mobile and desktop profiles.",
    "No console errors, broken links, accidental horizontal overflow, missing metadata or invalid structured data remain.",
    "Security headers, CSP reporting, production bundle contents and privacy language are verified.",
    "Latest-two desktop browser matrix and current mobile Safari/Chrome smoke tests pass.",
])
doc.add_heading("20.3 Release authority", level=2)
para("A technically green build is not sufficient approval. Production deployment requires an approved preview, explicit confirmation that the visual character has been preserved, and completion of all P0 requirements and hard release gates.")

doc.add_heading("21. Decision log", level=1)
table(["Decision", "Rationale", "Status"], [
    ("Restored production site is visual authority", "It represents the intentional character the owner chose after reverting the previous build.", "Locked"),
    ("Thesis is supporting, not visually overriding", "It supplies requirements and principles but cannot supersede the approved visual baseline.", "Locked"),
    ("Proof sheet contributes mechanisms, not a skin", "Its production language is useful selectively; its single-page/locked architecture is not.", "Locked"),
    ("TEDx RSU launches first", "One complete flagship case study creates more trust than several partial projects.", "Approved direction"),
    ("CV lives in About and Index", "It supports professional evaluation without becoming primary navigation clutter.", "Approved direction"),
    ("No database/CMS/large analytics stack in V1", "The release should solve proof and conversion with minimal operational weight.", "Locked"),
], [2.25, 3.55, 0.9])

doc.add_heading("22. Open inputs—not open product questions", level=1)
para("The product direction is sufficiently defined to begin. The remaining items are content or operational inputs, not requests to reconsider the site’s visual identity or architecture.")
bullets([
    "Final TEDx RSU media, captions, individual role, collaborators, dates and verifiable outcomes.",
    "Approved Explorations and Community entries, or approval of the drafted truthful empty states.",
    "Current CV PDF and confirmation of the five service descriptions.",
    "Working-principles copy if that optional homepage module is to ship.",
    "Vercel environment ownership, Resend sender/domain setup and shared rate-limit service choice.",
    "Final preview approval after the defined browser, accessibility and performance evidence is attached.",
])

doc.add_heading("23. Definition of done", level=1)
para("The release is done when a first-time design leader can arrive on the homepage, understand the positioning, inspect a complete and truthful flagship case study, verify relevant experience, obtain the CV and send a reliable enquiry—on keyboard, touch, reduced-motion and no-JavaScript paths—without encountering a broken route, unsupported claim or visual departure from the restored site.")

doc.core_properties.title = "Miebaka Portfolio Product Requirements Document"
doc.core_properties.subject = "Implementation-ready PRD for the Miebaka portfolio"
doc.core_properties.author = "Miebaka Joshua"
doc.core_properties.keywords = "portfolio, PRD, design engineering, case study, accessibility, performance"
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
