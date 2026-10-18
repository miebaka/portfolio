from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path('/Users/mac/Desktop/Career/Portfolio')
SPEC = ROOT / 'docs/superpowers/specs/2026-10-13-portfolio-design-engineering-review.md'
PLAN = ROOT / 'docs/superpowers/plans/2026-10-13-portfolio-design-engineering-implementation.md'
OUT = ROOT / 'output/portfolio-engineering-review/Portfolio Design Engineering Review.docx'


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for name, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{name}'))
        if node is None:
            node = OxmlElement(f'w:{name}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_border(cell, color='D9D9D9', size='4'):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = f'w:{edge}'
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), size)
        el.set(qn('w:color'), color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement('w:fldChar')
    fld_char1.set(qn('w:fldCharType'), 'begin')
    instr_text = OxmlElement('w:instrText')
    instr_text.set(qn('xml:space'), 'preserve')
    instr_text.text = ' PAGE '
    fld_char2 = OxmlElement('w:fldChar')
    fld_char2.set(qn('w:fldCharType'), 'end')
    run._r.extend([fld_char1, instr_text, fld_char2])


def add_rich_text(paragraph, text, size=None, color=None):
    parts = re.split(r'(`[^`]+`|\*\*[^*]+\*\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Courier New'
            run.font.size = Pt(9.5 if size is None else size)
        elif part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(part)
        if size is not None:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = RGBColor(*color)


def add_markdown(doc, path, skip_first_heading=True):
    lines = path.read_text(encoding='utf-8').splitlines()
    in_code = False
    first_heading_skipped = False
    for raw in lines:
        line = raw.rstrip()
        if line.startswith('```'):
            in_code = not in_code
            continue
        if in_code:
            p = doc.add_paragraph(style='Code')
            p.add_run(line)
            continue
        if not line.strip() or line.strip() == '---':
            continue
        if line.startswith('# '):
            if skip_first_heading and not first_heading_skipped:
                first_heading_skipped = True
                continue
            doc.add_heading(line[2:].strip(), level=1)
            continue
        if line.startswith('## '):
            doc.add_heading(line[3:].strip(), level=1)
            continue
        if line.startswith('### '):
            doc.add_heading(line[4:].strip(), level=2)
            continue
        if line.startswith('#### '):
            doc.add_heading(line[5:].strip(), level=3)
            continue
        if line.startswith('> '):
            p = doc.add_paragraph(style='Intense Quote')
            add_rich_text(p, line[2:].strip())
            continue
        checklist = re.match(r'^- \[([ xX])\] (.+)$', line)
        if checklist:
            p = doc.add_paragraph(style='Checklist')
            p.add_run('☐ ' if checklist.group(1) == ' ' else '☑ ')
            add_rich_text(p, checklist.group(2))
            continue
        bullet = re.match(r'^- (.+)$', line)
        if bullet:
            p = doc.add_paragraph(style='List Bullet')
            add_rich_text(p, bullet.group(1))
            continue
        numbered = re.match(r'^(\d+)\. (.+)$', line)
        if numbered:
            p = doc.add_paragraph(style='List Number')
            add_rich_text(p, numbered.group(2))
            continue
        p = doc.add_paragraph()
        add_rich_text(p, line)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.78)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(28, 32, 43)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.16

title = styles['Title']
title.font.name = 'Arial'
title.font.size = Pt(31)
title.font.bold = True
title.font.color.rgb = RGBColor(0, 0, 0)
title.paragraph_format.space_after = Pt(12)

for style_name, size, before, after in [
    ('Heading 1', 19, 18, 7),
    ('Heading 2', 14, 13, 5),
    ('Heading 3', 11.5, 10, 4),
]:
    style = styles[style_name]
    style.font.name = 'Arial'
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

for name in ('List Bullet', 'List Number'):
    style = styles[name]
    style.font.name = 'Arial'
    style.font.size = Pt(10.3)
    style.paragraph_format.space_after = Pt(3.5)

checklist = styles.add_style('Checklist', WD_STYLE_TYPE.PARAGRAPH)
checklist.font.name = 'Arial'
checklist.font.size = Pt(10)
checklist.paragraph_format.left_indent = Inches(0.18)
checklist.paragraph_format.first_line_indent = Inches(-0.18)
checklist.paragraph_format.space_after = Pt(3)

code = styles.add_style('Code', WD_STYLE_TYPE.PARAGRAPH)
code.font.name = 'Courier New'
code.font.size = Pt(9)
code.font.color.rgb = RGBColor(40, 44, 54)
code.paragraph_format.left_indent = Inches(0.18)
code.paragraph_format.space_after = Pt(3)

quote = styles['Intense Quote']
quote.font.name = 'Arial'
quote.font.size = Pt(10.5)
quote.font.italic = True
quote.font.color.rgb = RGBColor(28, 32, 43)

cover = doc.add_paragraph(style='Title')
cover.add_run('Portfolio Design Engineering Review')
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(24)
add_rich_text(subtitle, 'A structural improvement specification and implementation workflow for the Miebaka Joshua portfolio', size=14, color=(45, 52, 70))

summary = doc.add_paragraph()
summary.paragraph_format.space_after = Pt(18)
add_rich_text(summary, 'Main decision. ', size=11)
summary.runs[0].bold = True
add_rich_text(summary, 'The restored production website is the visual authority. The work below strengthens proof, navigation, accessibility, conversion and engineering quality without reviving the rejected redesign.', size=11)

meta = doc.add_table(rows=4, cols=2)
meta.autofit = False
meta.columns[0].width = Inches(1.55)
meta.columns[1].width = Inches(5.15)
for row in meta.rows:
    row.cells[0].width = Inches(1.55)
    row.cells[1].width = Inches(5.15)
for label, value, row in [
    ('Date', '13 October 2026', 0),
    ('Production baseline', 'Commit 54dcc41 and miebaka-portfolio.vercel.app', 1),
    ('Repository', 'github.com/miebaka/portfolio', 2),
    ('Scope', 'Design engineering, frontend, backend, accessibility, performance, security and delivery', 3),
]:
    meta.cell(row, 0).text = label
    meta.cell(row, 1).text = value
for row_i, row in enumerate(meta.rows):
    for col_i, cell in enumerate(row.cells):
        set_cell_margins(cell)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if col_i == 0:
            shade_cell(cell, 'E8EEFF')
            for run in cell.paragraphs[0].runs:
                run.bold = True
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.name = 'Arial'
                run.font.size = Pt(9.5)

doc.add_paragraph()
source_note = doc.add_paragraph()
add_rich_text(source_note, 'Evidence base. ', size=9.5)
source_note.runs[0].bold = True
add_rich_text(source_note, 'Live-site inspection, repository and commit analysis, both supplied Markdown documents, automated tests, and independent frontend and backend specialist audits.', size=9.5, color=(70, 75, 88))

doc.add_page_break()
add_markdown(doc, SPEC)

doc.add_section(WD_SECTION.NEW_PAGE)
doc.add_heading('Implementation workflow', level=1)
intro = doc.add_paragraph()
add_rich_text(intro, 'This plan is deliberately sequenced to protect the restored visual baseline. Each task creates an independently reviewable result and ends with verification before the next structural layer begins.')
add_markdown(doc, PLAN)

for sec in doc.sections:
    sec.footer.is_linked_to_previous = False
    footer = sec.footer.paragraphs[0]
    footer.clear()
    footer_run = footer.add_run('Miebaka Joshua portfolio engineering review   ')
    footer_run.font.name = 'Arial'
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(100, 104, 114)
    add_page_number(footer)

doc.core_properties.title = 'Portfolio Design Engineering Review'
doc.core_properties.subject = 'Design engineering audit, specification and implementation workflow'
doc.core_properties.author = 'Miebaka Joshua'
doc.core_properties.comments = 'Prepared from the production portfolio, repository, supplied design thesis and owner answers.'
doc.save(OUT)
print(OUT)
